"""Build a manual public catalog snapshot from two local Siigo exports.

Requires openpyxl. Workbooks are opened read-only; no credentials or API calls.
Run --help for inputs. Defaults to preparing/auditing outside the repository;
--write-catalog explicitly writes the public catalog and its source manifest.
The price-list and tax-basis flags must match what was verified in Siigo UI.
Keep the private output directory and source workbooks outside the repository.
"""

import argparse
import hashlib
import json
import math
import re
import shutil
import unicodedata
import warnings
from collections import Counter
from pathlib import Path

import openpyxl


CATEGORIES = {
    "Cerrajería", "Ferretería General", "Herramientas", "Herrajes para Muebles",
    "Tornillería y Fijación", "Adhesivos y Sellantes", "Eléctrico", "Fontanería",
    "Seguridad Industrial",
}
KNOWN_BRANDS = {
    "YALE", "STANLEY", "DEWALT", "MAKITA", "BOSCH", "TRUPER", "PRETUL", "TOTAL",
    "KL", "MHA", "IRWIN", "QUALITA", "FLEXON", "BAHCO", "VERA", "HERMEX",
    "BULDORK", "EINHELL", "TRAMONTINA", "3M",
}
PRODUCT_HEADERS = ["Tipo", "Código", "Nombre", "Unidad", "Precios", "Impuestos", "Stock", "Estado"]
SALES_HEADERS = [
    "Código producto", "Nombre producto", "Referencia fábrica", "Grupo inventario",
    "Cantidad vendida", "Valor bruto", "Descuento", "Subtotal", "Impuesto cargo",
    "Impuesto retención", "Total",
]
GROUP_CATEGORY = {}
for category, groups in {
    "Cerrajería": ["Cerraduras", "Candados", "Pasador", "Bisagras", "Cierre de Puerta"],
    "Herramientas": ["Herramienta", "Herramientas Eléctricas", "Abrasivos", "Grata", "Maquinaria", "maquinaria y equipo"],
    "Herrajes para Muebles": ["Herraje", "Cocinas", "Corredera", "Rodachina", "Riel", "Manijas", "Oficina"],
    "Tornillería y Fijación": ["Tornillo", "Puntilla", "Fijadores", "Grapas", "Abrazaderas", "Cadena/Cáncamo"],
    "Adhesivos y Sellantes": ["Pegante", "Cinta"],
    "Eléctrico": ["Eléctricos"],
    "Fontanería": ["Grifería"],
    "Seguridad Industrial": ["Seguridad Industrial"],
}.items():
    for group in groups:
        GROUP_CATEGORY[group] = category


def normalized(value):
    return " ".join(unicodedata.normalize("NFD", str(value or "")).encode("ascii", "ignore").decode().upper().split())


GROUP_CATEGORY = {normalized(group): category for group, category in GROUP_CATEGORY.items()}


def number(value, label):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Invalid numeric {label}: {value!r}")
    return int(value) if value == int(value) else value


def code(cell):
    # String cells are authoritative, including their leading zeroes.
    if isinstance(cell.value, str):
        return cell.value
    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) and float(cell.value).is_integer():
        value = str(int(cell.value))
        if re.fullmatch(r"0+", cell.number_format):
            return value.zfill(len(cell.number_format))
        raise ValueError("Numeric SKU without an explicit zero format requires review; refusing to infer its string representation.")
    raise ValueError(f"Invalid SKU in {cell.coordinate}")


def digest(path):
    checksum = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            checksum.update(block)
    return checksum.hexdigest()


def read_products(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Product export must have one worksheet.")
        sheet = workbook.worksheets[0]
        records = {}
        for row_number, cells in enumerate(sheet.iter_rows(), 1):
            values = [cell.value for cell in cells]
            if row_number == 4 and values[:8] != PRODUCT_HEADERS:
                raise ValueError(f"Unexpected product headers: {values[:8]}")
            if row_number <= 4 or not any(value is not None for value in values):
                continue
            if any(value is not None for value in values[8:]):
                raise ValueError("Unexpected additional product columns.")
            kind, _, name, unit, price, tax, stock, state = values[:8]
            sku = code(cells[1])
            if sku in records:
                raise ValueError(f"Duplicate product SKU: {sku}")
            if state not in {"Active", "Inactive"} or kind not in {"Producto", "Servicio"} or not isinstance(name, str) or not name.strip():
                raise ValueError(f"Unknown product type/state/name for {sku}")
            tax_match = re.fullmatch(r"IVA\s+(0|5|19)%", str(tax or "").strip(), re.IGNORECASE)
            records[sku] = {
                "sourceRow": row_number, "sku": sku, "type": kind, "name": name,
                "unit": unit, "price": number(price, f"price for {sku}"),
                "taxRaw": tax, "taxRate": int(tax_match[1]) if tax_match else None,
                "stockRaw": number(stock, f"stock for {sku}"), "active": state == "Active",
            }
        return records, {"sheet": sheet.title, "headerRow": 4, "headers": PRODUCT_HEADERS}
    finally:
        workbook.close()


def read_sales(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Sales export must have one worksheet.")
        sheet = workbook.worksheets[0]
        records = {}
        range_label = sheet.cell(5, 1).value
        for row_number, cells in enumerate(sheet.iter_rows(), 1):
            values = [cell.value for cell in cells]
            if row_number == 7 and values[:11] != SALES_HEADERS:
                raise ValueError(f"Unexpected sales headers: {values[:11]}")
            if row_number <= 7 or not any(value is not None for value in values):
                continue
            if values[0] is None or values[0] == "Total General" or str(values[0]).startswith("Procesado en:"):
                continue  # footer totals are not product rows
            sku = code(cells[0])
            if sku in records:
                raise ValueError(f"Duplicate sales SKU requires review: {sku}")
            quantity = number(values[4], f"quantity sold for {sku}")
            records[sku] = {
                "sourceRow": row_number, "sku": sku, "name": values[1], "reference": values[2],
                "group": values[3], "quantitySold": quantity,
            }
        return records, {"sheet": sheet.title, "headerRow": 7, "headers": SALES_HEADERS, "rangeLabel": range_label}
    finally:
        workbook.close()


def classify(name, group, legacy):
    text = normalized(name)
    # Only unambiguous product terms override historical categories/group errors.
    name_rules = [
        ("Seguridad Industrial", r"\b(GUANTE|GAFA|CARETA|CASCO|RESPIRADOR)\b"),
        ("Eléctrico", r"\b(TOMACORRIENTE|MULTITOMA|BOMBILLO|INTERRUPTOR|PASACABLE|GROMMET)\b|\bTOMA CORRIENTE\b|\bCINTA AISLANTE\b"),
        ("Fontanería", r"\b(GRIFO|GRIFERIA|DUCHA|SIFON)\b"),
        ("Cerrajería", r"\b(CERRADURA|CANDADO|CERROJO|CILINDRO|CIERRAPUERTA)\b"),
        ("Herramientas", r"\b(TALADRO|BROCA|LIJADORA|PULIDORA|MARTILLO|GRATA|SIERRA|DESTORNILLADOR|ATORNILLADOR)\b|\bDISCO (CORTE|LIJA|DIAMANTADO)\b|\bPISTOLA (SILICONA|PEGANTE)\b"),
        ("Adhesivos y Sellantes", r"\b(PEGANTE|SILICONA|SELLANTE|SIKAFLEX)\b"),
        ("Tornillería y Fijación", r"\b(TORNILLO|PUNTILLA|CHAZO|TARUGO|GRAPA|CANCAMO|TUERCA|ARANDELA)\b"),
        ("Herrajes para Muebles", r"\b(CORREDERA|RODACHINA|ALACENA|ENTREPANO|TIRADERA|CAZUELA|NIVELADOR|BOTON|MUEBLE)\b|\bBISAGRA .*\b(PARCHE|SEMIPARCHE|CODO|CAZOLETA)\b"),
    ]
    for category, pattern in name_rules:
        if re.search(pattern, text):
            return category, "explicit_name"
    if legacy.get("cat") in CATEGORIES:
        return legacy["cat"], "legacy"
    mapped = GROUP_CATEGORY.get(normalized(group))
    return (mapped, "source_group") if mapped else ("Ferretería General", "uncertain_group")


def brand_for(name, legacy):
    text = normalized(name)
    matches = [brand for brand in sorted(KNOWN_BRANDS) if re.search(r"(?<![A-Z0-9])" + re.escape(brand) + r"(?![A-Z0-9])", text)]
    if len(matches) == 1:
        return matches[0], "explicit_name"
    previous = legacy.get("brand")
    if previous in KNOWN_BRANDS and (not matches or previous in matches):
        return previous, "legacy"
    return "Sin marca", "unconfirmed"


def save_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--products", type=Path, required=True, help="Current Siigo product/service XLSX export")
    parser.add_argument("--sales", type=Path, required=True, help="Sales-by-product XLSX with credit notes excluded")
    parser.add_argument("--legacy-catalog", type=Path, required=True, help="Prior catalog JSON, used only for known category/brand metadata")
    parser.add_argument("--private-output", type=Path, required=True, help="Directory outside repository for audit/source copies")
    parser.add_argument("--date", required=True, help="Snapshot ISO date, verified in Siigo")
    parser.add_argument("--sales-from", required=True, help="Verified lower bound; does not assert the business started Siigo then")
    parser.add_argument("--sales-to", required=True)
    parser.add_argument("--price-list", required=True, help="Price list verified in Siigo UI")
    parser.add_argument("--prices-exclude-tax", action="store_true", required=True, help="Confirms the verified export price excludes IVA")
    parser.add_argument("--credit-notes-excluded", action="store_true", required=True, help="Confirms the verified sales UI filter")
    parser.add_argument("--expected-count", type=int, help="Fail if the public product count differs")
    parser.add_argument("--write-catalog", type=Path, help="Explicitly write public catalog JSON; otherwise prepare only")
    parser.add_argument("--manifest", type=Path, help="Public manifest path, required with --write-catalog")
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[1]
    private = args.private_output.resolve()
    if private == repo or repo in private.parents:
        raise ValueError("Private source/audit output must be outside the repository.")
    if args.write_catalog and not args.manifest:
        raise ValueError("--manifest is required when writing the catalog.")
    for value in (args.date, args.sales_from, args.sales_to):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Dates must use ISO YYYY-MM-DD.")
    warnings.simplefilter("ignore", UserWarning)
    private.mkdir(parents=True, exist_ok=True)
    sources = {}
    for label, path in (("products", args.products), ("sales", args.sales), ("legacyCatalog", args.legacy_catalog)):
        original_hash = digest(path)
        source_copy = private / f"source-{label}-{args.date}{path.suffix}"
        if source_copy.resolve() != path.resolve():
            if source_copy.exists() and digest(source_copy) != original_hash:
                raise ValueError(f"Existing source copy differs: {source_copy}")
            shutil.copy2(path, source_copy)
        if digest(source_copy) != original_hash:
            raise ValueError(f"Source copy checksum mismatch: {label}")
        sources[label] = {"original": str(path.resolve()), "copy": str(source_copy), "sha256": original_hash}
    products, product_meta = read_products(args.products)
    sales, sales_meta = read_sales(args.sales)
    legacy_items = json.loads(args.legacy_catalog.read_text(encoding="utf-8-sig"))
    legacy = {item["id"]: item for item in legacy_items}
    sold = {sku: row for sku, row in sales.items() if row["quantitySold"] > 0}
    missing = sorted(set(sold) - set(products))
    if missing:
        save_json(private / "missing-sold-skus.private.json", missing)
        raise ValueError(f"{len(missing)} sold SKUs missing from current export; review before publishing.")
    public = []
    joined = []
    excluded_services = []
    for sku, sale in sold.items():
        current = products[sku]
        if current["type"] == "Servicio":
            excluded_services.append({"sku": sku, "name": current["name"]})
            continue
        if not re.fullmatch(r"[a-zA-Z0-9._-]{1,120}", sku):
            raise ValueError(f"SKU needs an explicit route/API contract before publishing: {sku!r}")
        old = legacy.get(sku, {})
        category, category_basis = classify(current["name"], sale["group"], old)
        brand, brand_basis = brand_for(current["name"], old)
        verified = current["active"] and current["price"] > 0 and current["taxRate"] is not None
        item = {
            "id": sku, "nombre": current["name"].strip(), "precio": current["price"],
            "unidad": str(current["unit"]).strip() if current["unit"] else "Consultar unidad",
            "stock": max(0, current["stockRaw"]) if current["active"] else 0,
            "cat": category, "brand": brand, "priceVerified": verified,
        }
        if current["taxRate"] is not None:
            item["taxRate"] = current["taxRate"]
        if isinstance(sale["reference"], str) and sale["reference"].strip():
            item["ref"] = sale["reference"].strip()
        public.append(item)
        joined.append({"productSource": current, "salesSource": sale, "publicProduct": item,
                       "categoryBasis": category_basis, "brandBasis": brand_basis,
                       "legacyCategory": old.get("cat"), "legacyBrand": old.get("brand")})
    if args.expected_count is not None and len(public) != args.expected_count:
        raise ValueError(f"Expected {args.expected_count} public products; prepared {len(public)}.")
    assert len(public) == len({item["id"] for item in public})
    assert all(item["cat"] in CATEGORIES and item["stock"] >= 0 for item in public)
    assert all(item.get("taxRate") in {None, 0, 5, 19} for item in public)
    assert all("original" not in item and "disc" not in item for item in public)
    manifest = {
        "updatedAt": args.date, "productCount": len(public), "priceList": args.price_list,
        "pricesIncludeTax": False, "source": "Siigo",
        "salesRange": {"from": args.sales_from, "to": args.sales_to}, "includesCreditNotes": False,
        "manualSnapshot": True,
    }
    audit = {
        "sources": sources, "productExport": product_meta, "salesExport": sales_meta,
        "uiVerified": manifest,
        "counts": {"currentExport": len(products), "salesRows": len(sales), "soldPositive": len(sold),
                   "missingSoldSkus": len(missing), "excludedServices": len(excluded_services),
                   "publicProducts": len(public), "inactiveSold": sum(not row["productSource"]["active"] for row in joined),
                   "verifiedPrices": sum(item["priceVerified"] for item in public),
                   "consultPrices": sum(not item["priceVerified"] for item in public),
                   "zeroPrices": sum(item["precio"] == 0 for item in public),
                   "unknownTax": sum("taxRate" not in item for item in public),
                   "negativeStockClamped": sum(row["productSource"]["stockRaw"] < 0 for row in joined),
                   "blankUnits": sum(item["unidad"] == "Consultar unidad" for item in public),
                   "leadingZeroSkus": sum(bool(re.match(r"0[0-9]", item["id"])) for item in public),
                   "invalidRouteSkus": 0},
        "excludedServices": excluded_services,
        "categoryCounts": dict(Counter(item["cat"] for item in public)),
        "brandCounts": dict(Counter(item["brand"] for item in public)),
        "categoryBasisCounts": dict(Counter(row["categoryBasis"] for row in joined)),
        "categoryChanges": [{"sku": row["publicProduct"]["id"], "name": row["publicProduct"]["nombre"],
                             "before": row["legacyCategory"], "after": row["publicProduct"]["cat"], "basis": row["categoryBasis"]}
                            for row in joined if row["legacyCategory"] and row["legacyCategory"] != row["publicProduct"]["cat"]],
        "notes": ["Exact SKU join only; no leading-zero normalization.",
                  "Sales lower bound is the earliest date selectable in Siigo, not evidence of the actual adoption date.",
                  "Prices use the UI-verified list, exclude tax, and are a manual snapshot rather than an API integration.",
                  "Inactive sold items stay visible for consultation with stock zero and unverified price.",
                  "Services are excluded from this hardware catalog; uncertain categories/brands use conservative fallbacks."],
    }
    save_json(private / "catalog-prepared.private.json", {"audit": audit, "records": joined})
    save_json(private / "catalog-audit.private.json", audit)
    save_json(private / "proposed-public-catalog.json", public)
    save_json(private / "proposed-catalog-source.json", manifest)
    if args.write_catalog:
        save_json(args.write_catalog, public)
        save_json(args.manifest, manifest)
    print(json.dumps({"written": bool(args.write_catalog), "counts": audit["counts"],
                      "categoryCounts": audit["categoryCounts"], "brandCounts": audit["brandCounts"]}, ensure_ascii=True))


if __name__ == "__main__":
    main()
