"""Prepare a compact catalog using annual activity, without editing public data.

Select the smallest prefix covering the requested percent of annual net subtotal
within each existing catalog category. The baseline categories and Product data
come from --source-catalog unchanged. Stock is never a selection criterion.
Outputs, sales figures and audit remain in a private directory outside the repo.
This measures annual commercial activity, not recurring sales or inventory turns.
"""

import argparse
import hashlib
import importlib.util
import json
import re
import sys
import warnings
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl

sys.dont_write_bytecode = True
_spec = importlib.util.spec_from_file_location("siigo_import_rules", Path(__file__).with_name("import-siigo.py"))
siigo = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(siigo)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def read_annual(path, from_date, to_date):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Annual workbook must have exactly one sheet.")
        sheet = workbook.worksheets[0]
        printed_range = sheet.cell(5, 1).value
        siigo.validate_sales_range(printed_range, from_date, to_date)
        result, footer_labels = {}, []
        headers_seen = False
        for number, cells in enumerate(sheet.iter_rows(), 1):
            values = [cell.value for cell in cells]
            if number == 7:
                if values[:11] != siigo.SALES_HEADERS:
                    raise ValueError("Unexpected annual sales headers.")
                headers_seen = True
            if number <= 7 or not values or values[0] is None:
                continue
            if values[0] == "Total General" or str(values[0]).startswith("Procesado en:"):
                footer_labels.append(values[0])
                continue
            sku = siigo.code(cells[0])
            if sku in result:
                raise ValueError(f"Duplicate annual SKU requires review: {sku}")
            quantity = siigo.number(values[4], f"annual net quantity for {sku}")
            subtotal = siigo.number(values[7], f"annual net subtotal for {sku}")
            if quantity > 0 and subtotal < 0:
                raise ValueError(f"Positive quantity with negative subtotal requires review: {sku}")
            result[sku] = {"sku": sku, "netQuantity": quantity, "netSubtotal": Decimal(str(subtotal)), "row": number}
        if not headers_seen:
            raise ValueError("Annual workbook is missing its header row.")
        processed_label = next((value for value in footer_labels if str(value).startswith("Procesado en:")), None)
        processed = None
        if processed_label:
            match = re.fullmatch(r"PROCESADO EN:\s*([A-Z]+)\s*(\d{1,2})\s+(\d{4})\s+(\d{2}:\d{2})", siigo.normalized(processed_label))
            if match:
                day_label = f"De {match[1]}{match[2]} {match[3]} a {match[1]}{match[2]} {match[3]}"
                processed = siigo.parse_sales_range_label(day_label)[0] + "T" + match[4]
        return result, {"sheet": sheet.title, "printedRange": printed_range, "processedAt": processed,
                        "closingVerified": False, "footerLabels": footer_labels}
    finally:
        workbook.close()


def select(source, annual, coverage):
    groups = defaultdict(list)
    source_by_id = {item["id"]: item for item in source}
    if len(source_by_id) != len(source) or any(not isinstance(item["id"], str) for item in source):
        raise ValueError("Baseline catalog IDs must be unique strings.")
    if any(item["cat"] not in siigo.CATEGORIES for item in source):
        raise ValueError("Unknown baseline category requires review.")
    for sku, sale in annual.items():
        if sale["netQuantity"] > 0 and sku in source_by_id:
            groups[source_by_id[sku]["cat"]].append(sale)
    selected_ids, rationale, category_summary = set(), [], {}
    for category in sorted(groups):
        # Stable tie-break reproduces the reviewed proposal: subtotal desc, SKU desc.
        ranked = sorted(groups[category], key=lambda row: (row["netSubtotal"], row["sku"]), reverse=True)
        total = sum((row["netSubtotal"] for row in ranked), Decimal(0))
        if total <= 0:
            raise ValueError(f"Category has no positive annual subtotal: {category}")
        cutoff = total * coverage / Decimal(100)
        cumulative = Decimal(0)
        kept = []
        for rank, row in enumerate(ranked, 1):
            selected_ids.add(row["sku"])
            kept.append(row["sku"])
            cumulative += row["netSubtotal"]
            rationale.append({"sku": row["sku"], "categoryBaseline": category, "annualRankWithinCategory": rank,
                              "annualNetQuantity": row["netQuantity"], "annualNetSubtotal": row["netSubtotal"],
                              "reason": "selected_before_category_coverage_cutoff", "stockUsedForSelection": False})
            if cumulative >= cutoff:
                break
        category_summary[category] = {"recentPositiveReferences": len(ranked), "selectedReferences": len(kept),
                                      "netSubtotalAll": total, "netSubtotalSelected": cumulative,
                                      "coveragePercentActual": cumulative / total * 100,
                                      "selectedWithoutStock": sum(source_by_id[sku]["stock"] <= 0 for sku in kept)}
    chosen = [item for item in source if item["id"] in selected_ids]
    assert len(chosen) == len(selected_ids)
    assert all(item == source_by_id[item["id"]] for item in chosen)
    return chosen, rationale, category_summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-catalog", type=Path, required=True)
    parser.add_argument("--source-manifest", type=Path, required=True)
    parser.add_argument("--annual-sales", type=Path, required=True)
    parser.add_argument("--sales-from", required=True)
    parser.add_argument("--sales-to", required=True)
    parser.add_argument("--includes-credit-notes", action="store_true", required=True, help="Explicitly confirms the verified annual report setting")
    parser.add_argument("--coverage", type=Decimal, default=Decimal(95))
    parser.add_argument("--expected-source-count", type=int)
    parser.add_argument("--expected-selected-count", type=int)
    parser.add_argument("--private-output", type=Path, required=True)
    args = parser.parse_args()
    if not Decimal(0) < args.coverage <= Decimal(100):
        raise ValueError("Coverage must be greater than zero and at most 100.")
    repo = Path(__file__).resolve().parents[1]
    output = args.private_output.resolve()
    if output == repo or repo in output.parents:
        raise ValueError("Selection proposals and audits must remain outside the repository.")
    warnings.simplefilter("ignore", UserWarning)
    source_hash = sha(args.source_catalog)
    source_manifest_hash = sha(args.source_manifest)
    source = json.loads(args.source_catalog.read_text(encoding="utf-8-sig"))
    manifest = json.loads(args.source_manifest.read_text(encoding="utf-8-sig"))
    if args.expected_source_count is not None and len(source) != args.expected_source_count:
        raise ValueError("Baseline count differs from the reviewed category baseline.")
    if manifest["productCount"] != len(source):
        raise ValueError("Baseline manifest count does not match source catalog.")
    annual, annual_metadata = read_annual(args.annual_sales, args.sales_from, args.sales_to)
    chosen, rationale, category_summary = select(source, annual, args.coverage)
    if args.expected_selected_count is not None and len(chosen) != args.expected_selected_count:
        raise ValueError(f"Expected {args.expected_selected_count} selected products; prepared {len(chosen)}.")
    selection_manifest = {**manifest, "productCount": len(chosen), "salesRange": {"from": args.sales_from, "to": args.sales_to},
                          "includesCreditNotes": True,
                          "selection": {"method": "annual-net-subtotal-coverage-by-category", "coveragePercent": float(args.coverage),
                                        "categoryBaselineProductCount": len(source), "categoryBaseline": "existing-catalog-categories",
                                        "requiresPositiveAnnualNetQuantity": True, "usesStock": False, "frequencyVerified": False},
                          "salesReport": {"processedAt": annual_metadata["processedAt"], "closingVerified": False}}
    source_ids = {item["id"] for item in source}
    summary = {"sourceCatalogCount": len(source), "annualRows": len(annual),
               "recentPositiveCatalogReferences": sum(row["netQuantity"] > 0 and sku in source_ids for sku, row in annual.items()),
               "selectedCount": len(chosen), "selectedWithoutStock": sum(item["stock"] <= 0 for item in chosen),
               "selectedUnverifiedPrice": sum(item.get("priceVerified") is False for item in chosen),
               "selectedCategoryCounts": dict(Counter(item["cat"] for item in chosen)),
               "range": selection_manifest["salesRange"], "annualMetadata": annual_metadata,
               "categoryBaselineSha256": source_hash,
               "sources": {"sourceCatalog": str(args.source_catalog.resolve()), "sourceCatalogSha256": source_hash,
                           "sourceManifestSha256": source_manifest_hash, "annualSales": str(args.annual_sales.resolve()), "annualSalesSha256": sha(args.annual_sales)},
               "categoryAudit": category_summary,
               "notes": ["Annual net commercial activity; no claim of recurring sales or high inventory rotation.",
                         "Stock does not determine eligibility; sold out products remain selected.",
                         "The annual report was exported on the period's final day; complete closing is not verified.",
                         "Categories and all public Product fields are preserved exactly from the reviewed baseline."]}
    if sha(args.source_catalog) != source_hash or sha(args.source_manifest) != source_manifest_hash:
        raise ValueError("Source catalog/manifest changed during selection; rerun against a stable snapshot.")
    output.mkdir(parents=True, exist_ok=True)
    save(output / "selection-annual-95-ids.private.json", [item["id"] for item in chosen])
    save(output / "selection-annual-95-rationale.private.json", {"summary": summary, "selected": rationale})
    save(output / "proposed-compact-products.private.json", chosen)
    save(output / "proposed-compact-catalog-source.public.json", selection_manifest)
    print(json.dumps({key: summary[key] for key in ["sourceCatalogCount", "annualRows", "recentPositiveCatalogReferences", "selectedCount", "selectedWithoutStock", "selectedUnverifiedPrice", "selectedCategoryCounts"]}, ensure_ascii=True))


if __name__ == "__main__":
    main()
